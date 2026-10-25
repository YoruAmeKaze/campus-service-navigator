from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = "service-navigator-提交说明.docx"

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_border(cell, color="D9D9D9", size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)

def set_run_font(run, name="Aptos", size=10.5, bold=False, color="222222"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

def style_paragraph(paragraph, space_after=6, line=1.15):
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = line

def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    style_paragraph(p, 3, 1.1)
    set_run_font(p.add_run(text))
    return p

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
styles["Normal"].font.size = Pt(10.5)
for name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 11.5)):
    styles[name].font.name = "Aptos Display" if name == "Title" else "Aptos"
    styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles[name].font.size = Pt(size)
    styles[name].font.bold = True
    styles[name].font.color.rgb = RGBColor(0, 0, 0)

p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_run_font(p.add_run("Service Navigator 提交说明"), "Aptos Display", 22, True, "000000")
p = doc.add_paragraph()
style_paragraph(p, 14, 1.15)
set_run_font(p.add_run("通用机构服务导航技能的发布包、运行方式与验证记录"), size=11, color="666666")

doc.add_heading("提交结论", level=1)
p = doc.add_paragraph()
style_paragraph(p, 8, 1.2)
set_run_font(p.add_run("当前版本可提交。"), size=11, bold=True, color="1F4E79")
set_run_font(p.add_run("技能覆盖公共机构服务目录构建、办事入口导航和知识库导出三类场景。最新评审为 92.5 分，S 级；一票否决项全部通过。"))

doc.add_heading("交付文件", level=1)
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
headers = ["文件", "用途", "提交状态"]
for cell, text in zip(table.rows[0].cells, headers):
    set_cell_shading(cell, "1F4E79")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_border(cell)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run(text), size=10, bold=True, color="FFFFFF")
rows = [
    ("service-navigator.zip", "干净的技能发布包", "已生成并通过 staging 检查"),
    ("service-navigator-提交说明.docx", "本说明文档", "随发布包提交"),
]
for row in rows:
    cells = table.add_row().cells
    for i, text in enumerate(row):
        set_cell_border(cells[i])
        cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cells[i].paragraphs[0]
        style_paragraph(p, 2, 1.1)
        set_run_font(p.add_run(text), size=9.5)

doc.add_heading("核心能力", level=1)
for text in [
    "Build directory：从机构官方起始域名发现公共服务入口，记录来源、证据、状态和置信度。",
    "Navigate：依据目录和官方来源，为自然语言请求返回 1 至 3 个最佳办事入口。",
    "Export knowledge base：校验目录并生成 knowledge-base.jsonl、REPORT.md 和 service-directory.json。",
    "安全边界：只处理公共页面和服务元数据，不登录、不提交表单、不支付、不收集或保存个人敏感信息。",
]:
    add_bullet(doc, text)

doc.add_heading("发布包内容", level=1)
p = doc.add_paragraph()
style_paragraph(p, 5, 1.15)
set_run_font(p.add_run("ZIP 内包含运行所需文件："), bold=True)
set_run_font(p.add_run(" SKILL.md、LICENSE、agents/openai.yaml、assets/report-template.md、references/ 下的规范与示例，以及 scripts/ 下的三个运行脚本和打包检查脚本。"))
p = doc.add_paragraph()
style_paragraph(p, 8, 1.15)
set_run_font(p.add_run("发布包不包含："), bold=True)
set_run_font(p.add_run(" .git、tests、.temp、评审报告、编辑器目录和其他 ZIP 文件。"))

doc.add_heading("运行要求", level=1)
for text in [
    "Node.js 20 或更高版本。",
    "校验目录：node scripts/service-navigator.mjs validate path\\to\\service-directory.json",
    "导出知识库：node scripts/service-navigator.mjs export path\\to\\service-directory.json --out path\\to\\output",
    "省略 --out 时，输出到 .temp/service-output。输出路径必须位于当前工作目录内。",
]:
    add_bullet(doc, text)

doc.add_heading("验证记录", level=1)
table = doc.add_table(rows=1, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
for cell, text in zip(table.rows[0].cells, ["检查项", "结果"]):
    set_cell_shading(cell, "1F4E79")
    set_cell_border(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run(text), size=10, bold=True, color="FFFFFF")
checks = [
    ("node --test tests\\directory-validation.test.mjs", "5 项测试全部通过"),
    ("validate service-directory.example.json", "valid: true，serviceCount: 1"),
    ("export service-directory.example.json", "成功生成三件套输出"),
    ("check-package staging", "required 文件齐全，禁止路径为 0"),
    ("路径安全检查", "拒绝工作目录外的 --out 路径"),
]
for item, result in checks:
    cells = table.add_row().cells
    for i, text in enumerate((item, result)):
        set_cell_border(cells[i])
        cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cells[i].paragraphs[0]
        style_paragraph(p, 2, 1.1)
        set_run_font(p.add_run(text), size=9.5)

doc.add_heading("提交前检查", level=1)
for text in [
    "确认提交的是最新 service-navigator.zip。",
    "保留本说明文档，便于评审者了解运行方式和验证结果。",
    "若提交到 Git，确认 main 分支包含最新提交；网络推送失败时可稍后重试。",
]:
    add_bullet(doc, text)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_run_font(footer.add_run("Service Navigator | Submission Notes"), size=8, color="888888")

doc.save(OUT)
print(OUT)
